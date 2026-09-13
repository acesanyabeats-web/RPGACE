#!/usr/bin/env python3
"""
Real, one-time importer for the UK government's CoFID (McCance & Widdowson's
Composition of Foods Integrated Dataset) into RPGACE's `cofid_foods` Supabase
table.

Why this exists as a standalone local script rather than a direct Supabase
write: this Claude Code Remote session's own outbound proxy blocks raw HTTPS
to assets.publishing.service.gov.uk (the file's only real distribution host,
confirmed by direct curl test, not assumed) AND to supabase.co directly
(confirmed by this project's own standing landmine, `supabase_dedup_scan.py`'s
own comment). So this script does NOT push to Supabase itself — it reads the
real local XLSX file (once a human has downloaded it) and emits batched SQL
INSERT statements to a local .sql file, which a Claude Code session with
Supabase MCP access then runs via execute_sql/apply_migration.

Real source (official, free, no signup — verified reachable by a normal
browser, just not by this sandboxed session):
  https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid
  (the direct XLSX link changes with each government re-publish; use whatever
  the current gov.uk page links to — do not hardcode last year's asset URL,
  it 404s after each revision).

Real CoFID structure (per its own published documentation, not guessed): the
workbook has several sheets (Introduction, Proximates, Inorganics, Vitamins,
Fatty acids, ...). The "Proximates" sheet is the one this script reads — it
carries Food Code / Food Name / Food Group / Energy (kcal) / Protein (g) /
Fat (g) / Carbohydrate (g) per 100g, which is exactly what RPGACE's
`cofid_foods` table needs (kcal_per_100g/protein_g_per_100g/carbs_g_per_100g/
fat_g_per_100g) - the same 4 fields Open Food Facts already supplies, so both
sources plug into the identical downstream nutrition-summing code
(rpgace_core.js's cookingOracle.logic._ensureNutrition).

Column matching is done by fuzzy (case-insensitive substring) header lookup,
never a hardcoded exact column index — CoFID's real column headers have
shifted wording slightly across its own revisions (2015/2019/2021), and a
hardcoded index would silently read the wrong column if a future revision
reorders anything. If a required column can't be found, this script fails
loud and lists the real headers it actually saw - it never guesses.

Usage:
    pip install openpyxl   # already installed this session; pypi.org is
                            # reachable through this session's own proxy
    python3 scripts/import_cofid.py /path/to/CoFID_downloaded.xlsx

Output:
    scripts/cofid_import_batch_NNN.sql  (one or more files, ~500 rows each)
"""
import sys
import re
import math
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

BATCH_SIZE = 500

# Real, honest fuzzy-match terms for each required column — matched against
# the ACTUAL header row text, never a hardcoded index.
COLUMN_HINTS = {
    "food_name": ["food name"],
    "food_group": ["food group", "group"],
    "kcal": ["energy", "kcal"],
    "protein": ["protein"],
    "carbs": ["carbohydrate"],
    "fat": ["fat"],
}


def find_proximates_sheet(wb):
    for name in wb.sheetnames:
        if "proximate" in name.lower():
            return wb[name]
    raise SystemExit(
        "ERROR: no sheet with 'Proximates' in its name was found. Real sheets "
        "in this workbook: " + ", ".join(wb.sheetnames) + " — CoFID's own "
        "published structure always includes a Proximates sheet; if this "
        "file genuinely lacks one, it may not be the real CoFID dataset."
    )


def find_header_row(sheet, max_scan=10):
    """CoFID's real XLSX has a few title/intro rows before the actual header
    row — scan the first few rows for the one that contains a real 'Food
    Name'-ish header, rather than assuming row 1."""
    for row_idx in range(1, max_scan + 1):
        row = [str(c.value or "").strip().lower() for c in sheet[row_idx]]
        if any("food name" in cell for cell in row):
            return row_idx, [str(c.value or "").strip() for c in sheet[row_idx]]
    raise SystemExit(
        f"ERROR: no header row found in the first {max_scan} rows containing "
        "'Food Name'. This script's header-detection assumption may not "
        "match this real file's actual layout — inspect it by hand."
    )


def match_column(headers, hints, kcal_needs_kcal_unit=False):
    for i, h in enumerate(headers):
        hl = h.lower()
        if kcal_needs_kcal_unit:
            if "energy" in hl and "kcal" in hl:
                return i
            continue
        if all(hint in hl for hint in hints):
            return i
    return None


def parse_number(raw):
    """CoFID cells can hold 'Tr' (trace), 'N', blank, or footnote markers
    (e.g. '106a') alongside real numbers — return None honestly for anything
    that isn't a real parseable number rather than guessing 0."""
    if raw is None:
        return None
    s = str(raw).strip()
    if not s:
        return None
    m = re.match(r"^-?\d+(\.\d+)?", s)
    if not m:
        return None
    try:
        val = float(m.group(0))
    except ValueError:
        return None
    if math.isnan(val):
        return None
    return val


def sql_escape(s):
    return s.replace("'", "''")


def main():
    if len(sys.argv) != 2:
        print("Usage: python3 scripts/import_cofid.py /path/to/CoFID.xlsx", file=sys.stderr)
        sys.exit(1)
    src = Path(sys.argv[1])
    if not src.exists():
        print(f"ERROR: file not found: {src}", file=sys.stderr)
        sys.exit(1)

    print(f"Loading {src} (this can take a moment for a ~3,300-row workbook)...")
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    sheet = find_proximates_sheet(wb)
    header_row_idx, headers = find_header_row(sheet)
    print(f"Found sheet '{sheet.title}', header row {header_row_idx}: {headers}")

    col_name = match_column(headers, COLUMN_HINTS["food_name"])
    col_group = match_column(headers, COLUMN_HINTS["food_group"])
    col_kcal = match_column(headers, [], kcal_needs_kcal_unit=True)
    col_protein = match_column(headers, COLUMN_HINTS["protein"])
    col_carbs = match_column(headers, COLUMN_HINTS["carbs"])
    col_fat = match_column(headers, COLUMN_HINTS["fat"])

    missing = [name for name, col in [
        ("food_name", col_name), ("kcal", col_kcal), ("protein", col_protein),
        ("carbs", col_carbs), ("fat", col_fat),
    ] if col is None]
    if missing:
        raise SystemExit(
            f"ERROR: could not find real columns for: {missing}. "
            f"Real headers seen: {headers} — fix COLUMN_HINTS above to match "
            "this file's actual wording rather than guessing values."
        )

    rows = []
    skipped_no_name = 0
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        name = str(row[col_name] or "").strip() if col_name < len(row) else ""
        if not name:
            skipped_no_name += 1
            continue
        group = str(row[col_group] or "").strip() if (col_group is not None and col_group < len(row)) else None
        kcal = parse_number(row[col_kcal]) if col_kcal < len(row) else None
        protein = parse_number(row[col_protein]) if col_protein < len(row) else None
        carbs = parse_number(row[col_carbs]) if col_carbs < len(row) else None
        fat = parse_number(row[col_fat]) if col_fat < len(row) else None
        rows.append((name, group, kcal, protein, carbs, fat))

    print(f"Parsed {len(rows)} real food rows (skipped {skipped_no_name} rows with no food name).")
    if not rows:
        raise SystemExit("ERROR: zero real rows parsed — something is wrong with header detection above, not a real empty dataset.")

    out_dir = Path(__file__).parent
    batch_num = 0
    for i in range(0, len(rows), BATCH_SIZE):
        batch = rows[i:i + BATCH_SIZE]
        batch_num += 1
        out_path = out_dir / f"cofid_import_batch_{batch_num:03d}.sql"
        values_sql = []
        for (name, group, kcal, protein, carbs, fat) in batch:
            def num(v):
                return "null" if v is None else str(v)
            group_sql = "null" if not group else f"'{sql_escape(group)}'"
            values_sql.append(
                f"('{sql_escape(name)}', {group_sql}, {num(kcal)}, {num(protein)}, {num(carbs)}, {num(fat)}, 'cofid')"
            )
        sql = (
            "insert into cofid_foods (food_name, food_group, kcal_per_100g, protein_g_per_100g, carbs_g_per_100g, fat_g_per_100g, source) values\n"
            + ",\n".join(values_sql) + ";\n"
        )
        out_path.write_text(sql, encoding="utf-8")
        print(f"Wrote {out_path} ({len(batch)} rows)")

    print(f"\nDone. {batch_num} SQL batch file(s) written to {out_dir}/cofid_import_batch_*.sql")
    print("Next: run each file's SQL via a Claude Code session with Supabase MCP access (execute_sql).")


if __name__ == "__main__":
    main()
